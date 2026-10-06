"""Exports: the Excel workbooks' formulas calculate to the pipeline's values, and the Power BI
files form a valid star schema.

Workbook formulas are evaluated with an independent calculation engine (tests/excel_engine.py),
because openpyxl only writes formulas. The first half runs offline on small hand-built data;
the tests marked `integration` run against the project database and skip without it.
"""

import re
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from src.analytics import comps
from src.exports import excel, powerbi
from tests.excel_engine import Calculated

FY25, FY26, AS_OF = date(2025, 3, 31), date(2026, 3, 31), date(2026, 10, 6)
ROW = excel.FIRST_DATA_ROW
EXPECTED_SHEETS = {
    "financial_summary.xlsx": ["Company Overview", "Financial Ratios", "Growth Analysis",
                               "Valuation"],
    "comparable_companies.xlsx": ["Peer Set", "Operating Metrics", "Capital Structure",
                                  "Valuation Multiples", "Peer Statistics"],
    "market_analysis.xlsx": ["Returns", "Risk", "Correlation"],
}


# ------------------------------------------------------------- synthetic data ----

def company(ticker, group="it", sector="IT", sector_type="non_financial", entity="company"):
    return {"ticker": ticker, "company_name": f"{ticker} Ltd", "entity_type": entity,
            "sector_type": sector_type, "peer_group": group, "sector_name": sector,
            "industry_name": "Industry"}


COMPANIES = pd.DataFrame([
    company("T"), company("A"), company("B"), company("C"), company("D"),
    company("P1", "pharma", "Pharma"), company("P2", "pharma", "Pharma"),
    company("P3", "pharma", "Pharma"),
    company("BANK", "banks", "Banks", "bank"),
    {**company("^IDX", None, None, None, "index")},
])


def metric(ticker, name, value, period_type="annual", period_end=FY26, **extra):
    return {"ticker": ticker, "metric_name": name, "period_type": period_type,
            "period_end_date": AS_OF if period_type in ("ttm", "point_in_time") else period_end,
            "value": value, "na_reason": None if value is not None else "N/A (reason)",
            "method": extra.get("method"), "unit": "x", "as_of_date": AS_OF,
            "input_fields": extra.get("input_fields", {}), "reporting_currency": "INR",
            "is_translated": False}


def build_metrics() -> pd.DataFrame:
    rows = []
    margins = {"T": 0.27, "A": 0.20, "B": 0.22, "C": 0.24, "D": 0.26, "P1": 0.30, "P2": 0.10,
               "P3": 0.20, "BANK": None}
    multiples = {"T": 30.0, "A": 10.0, "B": 20.0, "C": 30.0, "D": 40.0, "P1": 12.0, "P2": 18.0,
                 "P3": None, "BANK": None}
    debt = {"T": 0.10, "A": 0.10, "B": 0.30, "C": None, "D": 0.20, "P1": 0.5, "P2": 0.7,
            "P3": 0.9, "BANK": None}
    growth = {"T": 0.10, "A": 0.05, "B": 0.08, "C": 0.12, "D": -0.02, "P1": 0.03, "P2": 0.04,
              "P3": 0.06, "BANK": 0.07}
    for ticker in margins:
        rows += [metric(ticker, "ebitda_margin", margins[ticker]),
                 metric(ticker, "net_margin", (margins[ticker] or 0.3) / 2),
                 metric(ticker, "debt_to_equity", debt[ticker]),
                 metric(ticker, "revenue_growth", growth[ticker]),
                 metric(ticker, "roe", 0.15, method="average_balance"),
                 metric(ticker, "ev_ebitda", multiples[ticker], "ttm"),
                 metric(ticker, "pe_ratio", 20.0 if ticker != "BANK" else 12.0, "ttm",
                        method="ttm"),
                 metric(ticker, "market_cap", 5e11, "ttm"),
                 metric(ticker, "return_1y", growth[ticker], "point_in_time"),
                 metric(ticker, "volatility_1y", 0.25, "point_in_time"),
                 metric(ticker, "max_drawdown_1y", -0.2, "point_in_time",
                        input_fields={"peak_date": "2026-01-05", "trough_date": "2026-03-02"})]
    rows.append(metric("^IDX", "return_1y", 0.04, "point_in_time"))
    return pd.DataFrame(rows)


def build_statements() -> pd.DataFrame:
    rows = []
    revenue = {"T": (1000.0, 1100.0), "A": (2000.0, 2100.0), "B": (500.0, None),
               "C": (-10.0, 50.0), "D": (300.0, 294.0), "P1": (100.0, 103.0),
               "P2": (100.0, 104.0), "P3": (100.0, 106.0), "BANK": (100.0, 107.0)}
    for ticker, (previous, current) in revenue.items():
        for period_end, value in ((FY25, previous), (FY26, current)):
            for item, scale in (("revenue", 1e7), ("net_income", 1e6)):
                rows.append({"ticker": ticker, "line_item": item, "period_end_date": period_end,
                             "period_type": "annual",
                             "value": None if value is None else value * scale,
                             "original_value": None if value is None else value * scale,
                             "original_currency": "INR", "sector_type": "x"})
    return pd.DataFrame(rows)


def build_correlations() -> pd.DataFrame:
    tickers = ["T", "A", "B"]
    values = {("T", "A"): 0.6, ("T", "B"): 0.2, ("A", "B"): 0.4}
    rows = []
    for a in tickers:
        for b in tickers:
            rows.append({"ticker_a": a, "ticker_b": b,
                         "correlation": 1.0 if a == b else values.get((a, b), values.get((b, a))),
                         "n_observations": 700, "start_date": date(2023, 10, 9),
                         "end_date": AS_OF})
    return pd.DataFrame(rows)


@pytest.fixture(scope="module")
def data() -> excel.ExportData:
    return excel.ExportData(companies=COMPANIES, metrics=build_metrics(),
                            statements=build_statements(), correlations=build_correlations(),
                            sources="yfinance", as_of=AS_OF)


@pytest.fixture(scope="module")
def workbooks(data, tmp_path_factory):
    folder = tmp_path_factory.mktemp("excel")
    return {path.name: path for path in excel.export_excel(data, folder)}


def all_cells(path):
    for sheet in load_workbook(path):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    yield sheet.title, cell


def close(actual, expected) -> bool:
    if isinstance(actual, str) or isinstance(expected, str) or actual is None or expected is None:
        return actual == expected
    return abs(actual - expected) <= 1e-9 * max(1.0, abs(expected))


# ------------------------------------------------------------------ structure ----

def test_workbooks_have_the_required_sheets_and_headers(workbooks):
    assert set(workbooks) == set(EXPECTED_SHEETS)
    for name, path in workbooks.items():
        book = load_workbook(path)
        assert book.sheetnames == EXPECTED_SHEETS[name]
        for sheet in book:
            assert sheet["A2"].value.startswith("Source: yfinance"), (name, sheet.title)
            assert sheet["A3"].value == f"As of: {AS_OF}", (name, sheet.title)
            assert sheet["A4"].value.startswith("Units: "), (name, sheet.title)
            assert "not investment advice" in sheet["A5"].value, (name, sheet.title)


def test_formulas_are_live_and_compatible(workbooks):
    formulas = [cell.value for path in workbooks.values() for _, cell in all_cells(path)
                if isinstance(cell.value, str) and cell.value.startswith("=")]
    joined = "\n".join(formulas)
    assert len(formulas) > 300
    for function in ("MEDIAN(", "PERCENTILE.INC(", "AVERAGEIFS(", "INDEX(", "MATCH(", "COUNTIF("):
        assert function in joined, function
    assert "XLOOKUP" not in joined.upper()                 # INDEX/MATCH only, for compatibility
    assert not load_workbook(workbooks["comparable_companies.xlsx"])["Peer Statistics"]._pivots


def test_missing_values_are_written_as_na_never_zero(workbooks):
    book = load_workbook(workbooks["comparable_companies.xlsx"])
    operating = book["Operating Metrics"]
    header = [c.value for c in operating[excel.HEADER_ROW]]
    margin_column = header.index("EBITDA margin") + 1
    tickers = [operating.cell(r, 1).value for r in range(ROW, ROW + 9)]
    bank_row = ROW + tickers.index("BANK")
    assert operating.cell(bank_row, margin_column).value == "N/A"
    # a metric with no data for anyone is N/A in every row, and nothing is zero-filled
    gross = header.index("Gross margin") + 1
    assert {operating.cell(r, gross).value for r in range(ROW, ROW + 9)} == {"N/A"}
    values = [operating.cell(r, c).value for r in range(ROW, ROW + 9) for c in range(4, 12)]
    assert 0 not in values
    # capital-structure metrics have their own sheet; the bank is N/A there too
    capital = book["Capital Structure"]
    assert [c.value for c in capital[excel.HEADER_ROW]][3:6] == ["Debt/equity",
                                                                 "Net debt/EBITDA",
                                                                 "Current ratio"]
    assert capital.cell(bank_row, 4).value == "N/A"
    assert capital.cell(ROW, 4).value == 0.10                      # T's debt/equity


# ----------------------------------------------------------- calculated values ----

STAT_COLUMNS = {"target_value": "C", "n_peers": "D", "peer_min": "E", "peer_p25": "F",
                "peer_median": "G", "peer_mean": "H", "peer_p75": "I", "peer_max": "J"}


def expected_statistics(result_row: dict) -> dict:
    """What each Peer Statistics cell should show for one metric."""
    n = result_row["n_peers"]
    expected = {}
    for key, column in STAT_COLUMNS.items():
        value = result_row[key]
        if key == "target_value":
            expected[column] = "N/A" if value is None else value
        elif key == "n_peers":
            expected[column] = n
        elif n == 0:
            expected[column] = "N/A"
        elif key in ("peer_p25", "peer_mean", "peer_p75") and n < 4:
            expected[column] = "–"                      # suppressed for small peer groups
        else:
            expected[column] = value
    return expected


def compare_peer_statistics(calculated: Calculated, result: dict) -> list:
    by_metric = {row["metric_name"]: row for row in result["rows"]}
    mismatches = []
    for offset, name in enumerate(excel.STATISTIC_METRICS):
        row = ROW + offset
        for column, expected in expected_statistics(by_metric[name]).items():
            actual = calculated.value("Peer Statistics", f"{column}{row}")
            if not close(actual, expected):
                mismatches.append((name, column, actual, expected))
        # position: the same number as the app's label ("63rd percentile" / "2nd of 4")
        label, shown = by_metric[name]["position_label"], calculated.value(
            "Peer Statistics", f"M{row}")
        if (label is None) != (shown == "N/A") or (
                label and re.findall(r"\d+", label) != re.findall(r"\d+", shown)):
            mismatches.append((name, "M", shown, label))
    return mismatches


@pytest.mark.parametrize("target", ["T", "A", "P1", "BANK"])
def test_peer_statistics_formulas_match_the_pipeline(workbooks, data, target):
    calculated = Calculated(workbooks["comparable_companies.xlsx"], {("Peer Set", "B7"): target})
    result = comps.compare(data.metrics, data.companies, target)
    assert calculated.value("Peer Set", "B9") == len(result["peers"])
    assert compare_peer_statistics(calculated, result) == []


def test_peer_statistics_by_hand(workbooks):
    """Target T: EV/EBITDA peers are A 10, B 20, C 30, D 40 (T's own 30 is excluded)."""
    calculated = Calculated(workbooks["comparable_companies.xlsx"])
    row = ROW + (excel.STATISTIC_METRICS).index("ev_ebitda")

    def stat(column):
        return calculated.value("Peer Statistics", f"{column}{row}")

    assert (stat("C"), stat("D")) == (30.0, 4)
    assert (stat("E"), stat("F"), stat("G"), stat("H"), stat("I"), stat("J")) == \
        (10.0, 17.5, 25.0, 25.0, 32.5, 40.0)
    assert stat("K") == pytest.approx(0.20)                 # 30 / 25 - 1
    assert stat("M") == "percentile 63"                     # (2 below + half of 1 equal) / 4

    # debt/equity: C has no value, so three peers (0.10, 0.30, 0.20): quartiles and mean hidden
    row = ROW + (excel.STATISTIC_METRICS).index("debt_to_equity")
    assert stat("D") == 3 and stat("G") == pytest.approx(0.20)
    assert (stat("F"), stat("H"), stat("I")) == ("–", "–", "–")
    assert stat("M") == "rank 3 of 4"                       # T at 0.10: B and D are higher

    # a percentage metric reports the gap in points: 27% - median(20, 22, 24, 26) = 4 points
    row = ROW + (excel.STATISTIC_METRICS).index("ebitda_margin")
    assert stat("G") == pytest.approx(0.23) and stat("K") == pytest.approx(0.04)


def test_changing_the_target_recalculates_everything(workbooks):
    default = Calculated(workbooks["comparable_companies.xlsx"])
    changed = Calculated(workbooks["comparable_companies.xlsx"], {("Peer Set", "B7"): "P1"})
    row = ROW + (excel.STATISTIC_METRICS).index("ev_ebitda")
    assert default.value("Peer Set", "B8") == "it" and changed.value("Peer Set", "B8") == "pharma"
    assert changed.value("Peer Statistics", "B7") == "P1"
    # P1's peers P2 (18) and P3 (N/A): one peer with a value
    assert changed.value("Peer Statistics", f"C{row}") == 12.0
    assert changed.value("Peer Statistics", f"D{row}") == 1
    assert changed.value("Peer Statistics", f"G{row}") == 18.0
    assert changed.value("Peer Statistics", f"M{row}") == "rank 2 of 2"


def test_growth_formulas_and_reconciliation(workbooks):
    calculated = Calculated(workbooks["financial_summary.xlsx"])
    book = load_workbook(workbooks["financial_summary.xlsx"])["Growth Analysis"]
    header = [c.value for c in book[excel.HEADER_ROW]]
    growth = excel.get_column_letter(header.index("Growth FY2026 (formula)") + 1)
    reconcile = excel.get_column_letter(header.index("Formula minus stored") + 1)
    tickers = [book.cell(r, 1).value for r in range(ROW, ROW + 9)]

    def cell(ticker, column):
        return calculated.value("Growth Analysis", f"{column}{ROW + tickers.index(ticker)}")

    assert cell("T", growth) == pytest.approx(0.10)         # 1,100 / 1,000 - 1
    assert cell("D", growth) == pytest.approx(-0.02)        # 294 / 300 - 1
    assert cell("B", growth) == "N/A"                       # FY2026 revenue missing
    assert cell("C", growth) == "N/A"                       # negative base: no growth rate
    assert cell("T", reconcile) == pytest.approx(0.0, abs=1e-12)


def test_sector_average_and_universe_median_formulas(workbooks):
    calculated = Calculated(workbooks["financial_summary.xlsx"])
    book = load_workbook(workbooks["financial_summary.xlsx"])["Financial Ratios"]
    header = [c.value for c in book[excel.HEADER_ROW]]
    column = excel.get_column_letter(header.index("EBITDA margin") + 1)
    labels = {book.cell(r, 3).value: r for r in range(ROW + 9, ROW + 20)
              if book.cell(r, 3).value}
    # IT: (27 + 20 + 22 + 24 + 26) / 5 = 23.8%; Pharma: (30 + 10 + 20) / 3 = 20%; Banks: none
    assert calculated.value("Financial Ratios", f"{column}{labels['IT']}") == pytest.approx(0.238)
    assert calculated.value("Financial Ratios", f"{column}{labels['Pharma']}") == \
        pytest.approx(0.20)
    assert calculated.value("Financial Ratios", f"{column}{labels['Banks']}") == "N/A"
    # median of the 8 values 10, 20, 20, 22, 24, 26, 27, 30 = 23%
    assert calculated.value("Financial Ratios", f"{column}{labels['Universe median']}") == \
        pytest.approx(0.23)


def test_market_workbook_formulas(workbooks):
    calculated = Calculated(workbooks["market_analysis.xlsx"])
    returns = load_workbook(workbooks["market_analysis.xlsx"])["Returns"]
    header = [c.value for c in returns[excel.HEADER_ROW]]
    excess = excel.get_column_letter(header.index("1Y return minus benchmark (formula)") + 1)
    # T returned 10%, the benchmark 4%: 6 points
    assert calculated.value("Returns", f"{excess}{ROW}") == pytest.approx(0.06)
    # correlation: T's row is 1.0, 0.6, 0.2 -> mean with the others (0.6 + 0.2) / 2 = 0.4
    assert calculated.value("Correlation", f"E{ROW}") == pytest.approx(0.4)
    assert calculated.value("Correlation", f"E{ROW + 1}") == pytest.approx(0.5)   # A: 0.6, 0.4


# ------------------------------------------------------------------ star schema ----

def star():
    return {
        "dim_sector": pd.DataFrame({"sector_key": [0, 1]}),
        "dim_company": pd.DataFrame({"company_key": [1, 2], "sector_key": [1, 0]}),
        "dim_date": pd.DataFrame({"date_key": [20260331, 20261006]}),
        "fact_market_prices": pd.DataFrame({"company_key": [1, 2], "date_key": [20261006] * 2}),
        "fact_financials": pd.DataFrame({"company_key": [1], "date_key": [20260331],
                                         "period_type": ["annual"], "statement": ["income"],
                                         "line_item": ["revenue"]}),
        "fact_metrics": pd.DataFrame({"company_key": [1], "date_key": [20260331],
                                      "period_type": ["annual"], "metric_name": ["roe"]}),
        "fact_valuation": pd.DataFrame({"company_key": [1], "date_key": [20261006],
                                        "period_end_date_key": [20260331],
                                        "valuation_basis": ["current"]}),
    }


def test_star_schema_validation():
    assert powerbi.validate_star_schema(star()) == []

    orphan = star()
    orphan["fact_metrics"] = orphan["fact_metrics"].assign(company_key=[99])
    assert powerbi.validate_star_schema(orphan) == [
        "fact_metrics.company_key: 1 value(s) not in dim_company.company_key"]

    duplicated = star()
    duplicated["dim_company"] = pd.concat([duplicated["dim_company"]] * 2)
    assert any("duplicate row(s)" in p for p in powerbi.validate_star_schema(duplicated))

    missing_date = star()
    missing_date["fact_valuation"] = missing_date["fact_valuation"].assign(
        period_end_date_key=[20200101])
    assert any("fact_valuation.period_end_date_key" in p
               for p in powerbi.validate_star_schema(missing_date))

    incomplete = star()
    del incomplete["dim_date"]
    assert "dim_date: missing" in powerbi.validate_star_schema(incomplete)

    null_key = star()
    null_key["fact_market_prices"] = null_key["fact_market_prices"].assign(date_key=[None, 1])
    assert any("null in key" in p for p in powerbi.validate_star_schema(null_key))


def test_correlation_helper_table_validation():
    pairs = pd.DataFrame({"company_a_key": [1, 1, 2, 2], "company_b_key": [1, 2, 1, 2],
                          "window_label": ["3y"] * 4, "correlation": [1.0, 0.4, 0.4, 1.0]})
    tables = {**star(), "agg_return_correlation": pairs}
    assert powerbi.validate_helpers(tables) == []
    # it is a helper, not part of the star schema: the star validation does not require it
    assert "agg_return_correlation" not in powerbi.ORDER
    assert powerbi.validate_star_schema(star()) == []
    unknown = {**star(), "agg_return_correlation": pairs.assign(company_b_key=[1, 2, 1, 77])}
    assert powerbi.validate_helpers(unknown) == [
        "agg_return_correlation.company_b_key: 1 value(s) not in dim_company.company_key"]
    repeated = {**star(), "agg_return_correlation": pd.concat([pairs, pairs.iloc[[0]]])}
    assert any("duplicate" in p for p in powerbi.validate_helpers(repeated))
    assert "agg_return_correlation: missing" in powerbi.validate_helpers(star())


# ------------------------------------------------- against the project database ----

@pytest.fixture(scope="module")
def project_exports(tmp_path_factory):
    from src.database.connection import get_engine
    from src.exports.run import generate_exports
    try:
        engine = get_engine()
        with engine.connect() as conn:
            if not conn.execute(text("SELECT count(*) FROM core.metrics")).scalar():
                pytest.skip("project database has no metrics")
            folder = tmp_path_factory.mktemp("exports")
            written = generate_exports(conn, folder)
    except Exception as exc:
        if isinstance(exc, pytest.skip.Exception):
            raise
        pytest.skip(f"project database not reachable: {type(exc).__name__}")
    return engine, folder, written


@pytest.mark.integration
def test_powerbi_csvs_match_the_star_schema(project_exports):
    engine, folder, written = project_exports
    tables = {name: pd.read_csv(folder / "powerbi" / f"{name}.csv") for name in powerbi.ORDER}
    assert powerbi.validate_star_schema(tables) == []
    with engine.connect() as conn:
        for name, table in tables.items():
            in_database = conn.execute(text(f"SELECT count(*) FROM mart.{name}")).scalar()
            assert len(table) == in_database == written["powerbi_rows"][name], name
            columns = [r[0] for r in conn.execute(text(
                "SELECT column_name FROM information_schema.columns WHERE table_schema = 'mart' "
                "AND table_name = :t ORDER BY ordinal_position"), {"t": name})]
            assert list(table.columns) == columns, name
        prices = conn.execute(text("SELECT count(*) FROM core.market_prices")).scalar()
    assert len(tables["fact_market_prices"]) == prices
    dates = tables["dim_date"]
    assert dates["date_key"].is_monotonic_increasing and dates["date_key"].is_unique
    # the date dimension has no gaps: one row per calendar day
    days = pd.to_datetime(dates["date"])
    assert (days.diff().dropna() == pd.Timedelta(days=1)).all()
    april = dates[dates["date"] == "2025-04-01"].iloc[0]
    assert (april["fiscal_year_label"], april["fiscal_quarter_label"]) == ("FY2026", "Q1 FY2026")
    assert {"company", "index"} == set(tables["dim_company"]["entity_type"])

    # the correlation helper table: both orderings and the diagonal for each window
    pairs = pd.read_csv(folder / "powerbi" / "agg_return_correlation.csv")
    assert list(pairs.columns) == ["company_a_key", "company_b_key", "window_label",
                                   "window_start", "window_end", "correlation", "n_obs"]
    assert powerbi.validate_helpers({**tables, "agg_return_correlation": pairs}) == []
    assert len(pairs) == written["powerbi_rows"]["agg_return_correlation"]
    n_tickers = len(tables["dim_company"])
    assert (pairs.groupby("window_label").size() == n_tickers ** 2).all()
    diagonal = pairs[pairs["company_a_key"] == pairs["company_b_key"]]
    assert (diagonal["correlation"] == 1.0).all()
    mirrored = pairs.merge(pairs, left_on=["company_a_key", "company_b_key", "window_label"],
                           right_on=["company_b_key", "company_a_key", "window_label"])
    assert (mirrored["correlation_x"] - mirrored["correlation_y"]).abs().max() < 1e-12


@pytest.mark.integration
@pytest.mark.parametrize("target", ["TCS.NS", "HINDUNILVR.NS", "HDFCBANK.NS", "INFY.NS"])
def test_exported_comps_workbook_matches_stored_comparisons(project_exports, target):
    engine, folder, _ = project_exports
    from src.database import queries
    with engine.connect() as conn:
        metrics, companies = queries.load_metrics(conn), queries.load_companies(conn)
        stored = pd.read_sql(text("""
            SELECT p.metric_name, p.n_peers, p.peer_median, p.target_value
            FROM core.peer_comparisons p JOIN core.companies c USING (company_id)
            WHERE c.ticker = :t"""), conn, params={"t": target}).set_index("metric_name")
    calculated = Calculated(folder / "excel" / "comparable_companies.xlsx",
                            {("Peer Set", "B7"): target})
    assert compare_peer_statistics(calculated, comps.compare(metrics, companies, target)) == []
    # and against what the pipeline stored in PostgreSQL
    for offset, name in enumerate(excel.STATISTIC_METRICS):
        row = ROW + offset
        assert calculated.value("Peer Statistics", f"D{row}") == stored.at[name, "n_peers"]
        median = stored.at[name, "peer_median"]
        assert close(calculated.value("Peer Statistics", f"G{row}"),
                     "N/A" if pd.isna(median) else median), name


@pytest.mark.integration
def test_exported_growth_formulas_reconcile_with_stored_growth(project_exports):
    _, folder, _ = project_exports
    path = folder / "excel" / "financial_summary.xlsx"
    calculated = Calculated(path)
    sheet = load_workbook(path)["Growth Analysis"]
    header = [c.value for c in sheet[excel.HEADER_ROW]]
    column = excel.get_column_letter(header.index("Formula minus stored") + 1)
    differences, row = [], ROW
    while sheet.cell(row, 1).value:
        differences.append(calculated.value("Growth Analysis", f"{column}{row}"))
        row += 1
    numeric = [d for d in differences if not isinstance(d, str)]
    assert len(numeric) >= 20                                # most companies have a growth rate
    assert max(abs(d) for d in numeric) < 1e-9               # the workbook agrees with the pipeline
